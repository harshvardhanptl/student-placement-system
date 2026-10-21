"""
Resume Parser Module
Extracts, validates, and normalizes text from PDF and DOCX files.
Designed with security controls and graceful handling of malformed or scanned documents.
"""

import os
import re
import io
from typing import Dict, Any, Union
try:
    import pymupdf as fitz
except ImportError:
    import fitz
from docx import Document
import config


def validate_resume_file(uploaded_file: Any) -> Dict[str, Any]:
    """
    Validates file extension and size constraints.
    Accepts Streamlit UploadedFile or file path/buffer.
    """
    if uploaded_file is None:
        return {"valid": False, "error": "No file was provided."}

    file_name = getattr(uploaded_file, "name", str(uploaded_file))
    _, ext = os.path.splitext(file_name)
    ext = ext.lower()

    if ext not in config.ALLOWED_EXTENSIONS:
        return {
            "valid": False,
            "error": f"Unsupported file format '{ext}'. Please upload a PDF (.pdf) or Word document (.docx)."
        }

    # Size check
    file_size = getattr(uploaded_file, "size", None)
    if file_size is None and hasattr(uploaded_file, "getbuffer"):
        file_size = len(uploaded_file.getbuffer())
    elif file_size is None and isinstance(uploaded_file, (bytes, bytearray)):
        file_size = len(uploaded_file)
    elif file_size is None and isinstance(uploaded_file, str) and os.path.exists(uploaded_file):
        file_size = os.path.getsize(uploaded_file)

    if file_size and file_size > config.MAX_FILE_SIZE_BYTES:
        return {
            "valid": False,
            "error": f"File exceeds maximum allowed size of {config.MAX_FILE_SIZE_MB}MB ({file_size / (1024*1024):.2f}MB detected)."
        }

    return {"valid": True, "extension": ext, "filename": file_name, "size_bytes": file_size}


def sanitize_text(raw_text: str) -> str:
    """
    Sanitizes raw extracted text by removing control characters,
    normalizing whitespace, and trimming excessive newlines.
    """
    if not raw_text:
        return ""
    # Remove null bytes and non-printable control characters (except newline, carriage return, tab)
    sanitized = re.sub(r"[\x00-\x08\x0b\x0c\x0e-\x1f\x7f]", " ", raw_text)
    # Normalize multiple tabs/spaces
    sanitized = re.sub(r"[ \t]+", " ", sanitized)
    # Normalize excessive linebreaks (more than 2 to 2)
    sanitized = re.sub(r"\n\s*\n\s*\n+", "\n\n", sanitized)
    return sanitized.strip()


def extract_text_from_pdf(file_bytes: bytes) -> Dict[str, Any]:
    """
    Extracts text from PDF bytes using PyMuPDF (fitz).
    Detects scanned/image-only PDFs.
    """
    try:
        doc = fitz.open(stream=file_bytes, filetype="pdf")
        if doc.is_encrypted:
            return {
                "success": False,
                "text": "",
                "num_pages": len(doc),
                "error": "The PDF file is encrypted or password-protected. Please provide an unencrypted resume."
            }

        num_pages = len(doc)
        full_text = []

        for page_idx in range(num_pages):
            page = doc.load_page(page_idx)
            page_text = page.get_text("text")
            if page_text:
                full_text.append(page_text)

        doc.close()
        combined_text = sanitize_text("\n".join(full_text))

        # Check for scanned or image-based PDF
        warning = None
        if len(combined_text) < 50:
            warning = (
                "Notice: Very little readable text was extracted (< 50 characters). "
                "This document may be an image-only or scanned PDF. "
                "For accurate skill analysis, please upload a text-selectable PDF or DOCX file."
            )

        return {
            "success": True,
            "text": combined_text,
            "num_pages": num_pages,
            "warning": warning,
            "error": None
        }

    except Exception as e:
        return {
            "success": False,
            "text": "",
            "num_pages": 0,
            "error": f"Error parsing PDF document: {str(e)}"
        }


def extract_text_from_docx(file_bytes: bytes) -> Dict[str, Any]:
    """
    Extracts text from DOCX bytes using python-docx.
    """
    try:
        doc_io = io.BytesIO(file_bytes)
        document = Document(doc_io)
        full_text = []

        # Paragraphs
        for para in document.paragraphs:
            if para.text.strip():
                full_text.append(para.text.strip())

        # Tables (e.g., education, skills tables)
        for table in document.tables:
            for row in table.rows:
                row_text = " | ".join(cell.text.strip() for cell in row.cells if cell.text.strip())
                if row_text:
                    full_text.append(row_text)

        combined_text = sanitize_text("\n".join(full_text))

        warning = None
        if len(combined_text) < 50:
            warning = "Notice: Very little text extracted. The document may be empty or formatted with unsupported objects."

        return {
            "success": True,
            "text": combined_text,
            "num_pages": 1,
            "warning": warning,
            "error": None
        }

    except Exception as e:
        return {
            "success": False,
            "text": "",
            "num_pages": 0,
            "error": f"Error parsing DOCX document: {str(e)}"
        }


def parse_resume(uploaded_file: Any) -> Dict[str, Any]:
    """
    High-level facade to validate and parse resume documents.
    Accepts Streamlit UploadedFile, file-like object, or file path.
    """
    validation = validate_resume_file(uploaded_file)
    if not validation["valid"]:
        return {
            "success": False,
            "text": "",
            "filename": getattr(uploaded_file, "name", "unknown"),
            "error": validation["error"],
            "warning": None
        }

    # Read bytes
    try:
        if hasattr(uploaded_file, "read"):
            file_bytes = uploaded_file.read()
            # Reset pointer if possible
            if hasattr(uploaded_file, "seek"):
                uploaded_file.seek(0)
        elif isinstance(uploaded_file, (bytes, bytearray)):
            file_bytes = bytes(uploaded_file)
        elif isinstance(uploaded_file, str) and os.path.exists(uploaded_file):
            with open(uploaded_file, "rb") as f:
                file_bytes = f.read()
        else:
            return {"success": False, "text": "", "error": "Unable to read file stream."}
    except Exception as e:
        return {"success": False, "text": "", "error": f"Failed reading file content: {str(e)}"}

    ext = validation["extension"]
    if ext == ".pdf":
        result = extract_text_from_pdf(file_bytes)
    elif ext == ".docx":
        result = extract_text_from_docx(file_bytes)
    else:
        return {"success": False, "text": "", "error": f"Unsupported extension: {ext}"}

    result["filename"] = validation.get("filename", "resume")
    result["file_size_bytes"] = len(file_bytes)
    return result
